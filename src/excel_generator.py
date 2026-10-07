import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from typing import Dict, Any, Optional, List


def get_rr_rate_fill(rate: float, target_mp_rate: float = 0.005) -> Optional[PatternFill]:
    """
    Applies the exact 3-tier conditional color palette from the Keynote template:
    - >= 3.0%: Coral / Red (#FF9781) - 'No RC'
    - 2.0% to 3.0%: Orange (#FFD38A) - 'RC identified no fix available'
    - >= target_mp_rate (default 0.5%): Light Yellow (#FFFC98) - 'Validation ongoing'
    - < target_mp_rate: Normal white (no fill)
    """
    if rate >= 0.03:
        return PatternFill(start_color="FF9781", end_color="FF9781", fill_type="solid")
    elif rate >= 0.02:
        return PatternFill(start_color="FFD38A", end_color="FFD38A", fill_type="solid")
    elif rate >= target_mp_rate:
        return PatternFill(start_color="FFFC98", end_color="FFFC98", fill_type="solid")
    return None


def build_excel_workbook(
    report_data: Dict[str, Any],
    table_items_limit: Optional[int] = None,
    selected_stations: Optional[List[str]] = None,
    target_mp_rate: float = 0.005
) -> "openpyxl.Workbook":
    """
    Pure compute phase: builds and returns an openpyxl Workbook fully in memory.
    No file I/O is performed. Call wb.save(path) separately when ready to write.
    """
    wb = openpyxl.Workbook()

    ws = wb.active
    if table_items_limit == 5:
        ws.title = "RR Top5 issue"
    elif table_items_limit is not None and table_items_limit > 0:
        ws.title = f"RR Top{table_items_limit} issue"
    else:
        ws.title = "RR Retest issue"
    wb.create_sheet(title="Sheet2")

    col_widths = {
        'A': 5.34, 'B': 9.0, 'C': 18.16, 'D': 9.5, 'E': 13.0,
        'F': 10.84, 'G': 71.34, 'H': 9.34, 'I': 9.34,
        'J': 21.16, 'K': 81.66, 'L': 30.0, 'M': 41.34
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    ws.row_dimensions[1].height = 15.0
    ws.row_dimensions[2].height = 20.4
    ws.row_dimensions[3].height = 17.6

    font_title1 = Font(name="Calibri", size=11, bold=False)
    font_title2 = Font(name="Calibri", size=14, bold=True)
    font_header = Font(name="Calibri", size=12, bold=True)
    font_data  = Font(name="Calibri", size=11, bold=False)

    fill_header = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")

    thin_border_side = Side(border_style="thin", color="000000")
    thin_border = Border(
        left=thin_border_side, right=thin_border_side,
        top=thin_border_side, bottom=thin_border_side
    )

    align_center = Alignment(horizontal="center", vertical="center")
    align_left   = Alignment(horizontal="left",   vertical="center")
    align_sn     = Alignment(horizontal="left",   vertical="center", wrap_text=True)

    ws["A1"] = "Tool"
    ws["A1"].font = font_title1

    if table_items_limit == 5:
        ws["A2"] = "RR Top5 Issue"
    elif table_items_limit is not None and table_items_limit > 0:
        ws["A2"] = f"RR Top{table_items_limit} Issue"
    else:
        ws["A2"] = "RR Retest Issue"
    ws["A2"].font = font_title2

    for cell_id, text in [
        ("A3", "NO."), ("B3", " Product "), ("C3", " Station "),
        ("D3", "Input"), ("E3", "Retest"), ("F3", "RR Rate"),
        ("G3", "Top Issues"), ("H3", "Item Q'ty"), ("I3", "Item RR"),
        ("J3", "SN"), ("K3", "RR Symptoms"), ("L3", "Category"), ("M3", "FACA"),
    ]:
        cell = ws[cell_id]
        cell.value = text
        cell.font = font_header
        cell.fill = fill_header
        cell.border = thin_border
        cell.alignment = align_center

    current_row = 4
    st_counter = 1
    for station in report_data.get("stations", []):
        st_name = station["station"]
        if selected_stations and st_name not in selected_stations:
            continue

        start_row  = current_row
        st_no      = st_counter
        st_counter += 1
        st_prod    = station["product"]
        st_input   = station["input"]
        st_retest  = station["retest"]
        st_rate    = station["retest_pct"]

        issues = station.get("issues", [])
        if table_items_limit is not None and table_items_limit > 0:
            issues = issues[:table_items_limit]

        num_rows = max(len(issues), 1)
        end_row  = start_row + num_rows - 1

        for idx in range(num_rows):
            r = start_row + idx
            ws.row_dimensions[r].height = 40.0 if issues else 15.0
            for c in range(1, 14):   # 13 columns now (A-M)
                cell = ws.cell(row=r, column=c)
                cell.border = thin_border
                cell.font = font_data
                cell.alignment = align_center

            if idx < len(issues):
                iss = issues[idx]
                ws.cell(row=r, column=7, value=iss["issue_name"]).alignment = align_left
                ws.cell(row=r, column=8, value=iss["item_qty"]).alignment = align_center
                c9 = ws.cell(row=r, column=9, value=iss["item_rr"])
                c9.alignment = align_center
                c9.number_format = "0.00%"
                sn_val = iss.get("sn_str", "")
                ws.cell(row=r, column=10, value=sn_val if sn_val else None).alignment = align_sn
                ws.cell(row=r, column=11, value=iss["rr_symptom"]).alignment = align_left
                # Column 12 (L - Category): blank with border — user fills manually
                # Column 13 (M - FACA): blank with border — user fills manually

        ws.cell(row=start_row, column=1, value=st_no)
        ws.cell(row=start_row, column=2, value=st_prod)
        ws.cell(row=start_row, column=3, value=st_name)
        ws.cell(row=start_row, column=4, value=st_input)
        ws.cell(row=start_row, column=5, value=st_retest)
        c6 = ws.cell(row=start_row, column=6, value=st_rate)
        c6.number_format = "0.00%"

        rate_fill = get_rr_rate_fill(st_rate, target_mp_rate)
        if rate_fill:
            for r in range(start_row, end_row + 1):
                ws.cell(row=r, column=6).fill = rate_fill

        for col in range(1, 7):
            ws.cell(row=start_row, column=col).alignment = align_center

        if num_rows > 1:
            for col in range(1, 7):
                ws.merge_cells(
                    start_row=start_row, end_row=end_row,
                    start_column=col, end_column=col
                )

        current_row = end_row + 1

    return wb


def generate_excel_report(
    report_data: Dict[str, Any],
    output_path: str,
    table_items_limit: Optional[int] = None,
    selected_stations: Optional[List[str]] = None,
    target_mp_rate: float = 0.005
) -> str:
    """Compute + save in one call (backward-compatible wrapper)."""
    wb = build_excel_workbook(report_data, table_items_limit, selected_stations, target_mp_rate)
    wb.save(output_path)
    return output_path
